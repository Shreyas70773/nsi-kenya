import openpyxl, io
path = r"C:\Users\Shreyas Sunil\Documents\Claude\Projects\SEO + GEO Autopilot\North Star Impex — Quick-Win & AI Question Gap Map.xlsx"
out = io.StringIO()
wb = openpyxl.load_workbook(path, data_only=True)
out.write("SHEETS: " + str(wb.sheetnames) + "\n")
for ws in wb.worksheets:
    out.write("\n" + "=" * 80 + "\n")
    out.write(f"SHEET: {ws.title} ({ws.max_row} rows x {ws.max_column} cols)\n")
    out.write("=" * 80 + "\n")
    for row in ws.iter_rows(values_only=True):
        cells = [("" if c is None else str(c)) for c in row]
        if any(c.strip() for c in cells):
            out.write(" | ".join(cells).rstrip(" |") + "\n")
with open(r"C:\Users\Shreyas Sunil\northstar-website\scripts\xlsx-dump.txt", "w", encoding="utf-8") as f:
    f.write(out.getvalue())
print("written")
